"""Extracción y contraste de la planeación pedagógica GFPI-F-134."""

from __future__ import annotations

import math
import re
import unicodedata
from copy import deepcopy
from datetime import date, datetime, timedelta
from pathlib import Path
from difflib import SequenceMatcher
from xml.etree.ElementTree import iterparse

import openpyxl
from openpyxl.utils.cell import range_boundaries

from app.services.importacion_ficha import clasificar_competencia


_CAMPOS = {
    'denominacion del programa de formacion': 'nombre_programa',
    'codigo y version del programa de formacion': 'codigo_programa',
    'fecha de elaboracion': 'fecha_elaboracion',
}


def _texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return re.sub(r'\s+', ' ', str(valor).replace('\xa0', ' ').strip())


def _clave(valor):
    texto = unicodedata.normalize('NFKD', _texto(valor))
    return ''.join(c for c in texto if not unicodedata.combining(c)).lower()


def _numero(valor):
    if valor is None or _texto(valor) == '':
        return 0.0
    try:
        numero = float(str(valor).replace(',', '.'))
    except (TypeError, ValueError):
        return 0.0
    if not math.isfinite(numero):
        raise ValueError('La duración debe ser un número finito. Revise las horas del archivo de planeación.')
    return numero


def _fecha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = _texto(valor)
    for formato in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def _valor_combinado(hoja, mapa, fila, columna):
    return _texto(hoja[columna - 1])


def _mapa_combinadas(hoja):
    """Lee rangos compactos desde el XML; no expande cada celda combinada."""
    mapa = {}
    with hoja.parent._archive.open(hoja._worksheet_path) as xml:
        padres = []
        for evento, elemento in iterparse(xml, events=('start', 'end')):
            if evento == 'start':
                padres.append(elemento)
                continue
            if elemento.tag.endswith('}mergeCell'):
                min_col, min_row, max_col, max_row = range_boundaries(elemento.attrib['ref'])
                mapa.setdefault(min_row, []).append((min_col, max_col, max_row))
            if len(padres) > 1:
                padres[-2].remove(elemento)
            padres.pop()
            elemento.clear()
    return mapa


def _filas_combinadas(hoja, mapa):
    """Propaga únicamente los rangos activos durante la lectura secuencial."""
    activos = []
    for numero, fila in enumerate(hoja.iter_rows(max_col=17, values_only=True), start=1):
        valores = list(fila)
        activos = [rango for rango in activos if rango[2] >= numero]
        for min_col, max_col, max_row in mapa.pop(numero, []):
            if min_col <= 17:
                activos.append((min_col, min(max_col, 17), max_row, valores[min_col - 1]))
        for min_col, max_col, _, valor in activos:
            for columna in range(min_col - 1, max_col):
                valores[columna] = valor
        yield numero, valores


def _identificar_campo(clave: str) -> str | None:
    clave_limpia = clave.rstrip(':').strip()
    if not clave_limpia:
        return None
    if any(k in clave_limpia for k in ('codigo y version del programa', 'codigo y version de programa', 'codigo del programa', 'codigo programa')):
        return 'codigo_programa'
    if any(k in clave_limpia for k in ('denominacion del programa', 'nombre del programa', 'programa de formacion', 'denominacion programa')):
        return 'nombre_programa'
    if any(k in clave_limpia for k in ('fecha de elaboracion', 'fecha elaboracion', 'fecha elaboracion de la planeacion')):
        return 'fecha_elaboracion'
    if any(k in clave_limpia for k in ('codigo de la ficha', 'ficha de caracterizacion', 'numero de ficha', 'codigo ficha')):
        return 'codigo_ficha'
    return None


def _metadatos(hoja):
    resultado = {}
    for fila in hoja.iter_rows(max_row=25, values_only=True):
        valores = [_texto(valor) for valor in fila]
        for posicion, valor in enumerate(valores):
            clave = _clave(valor)
            campo = _identificar_campo(clave)
            if not campo or campo in resultado:
                continue
            siguiente = next((item for item in valores[posicion + 1:] if item), '')
            if campo == 'fecha_elaboracion':
                siguiente_crudo = fila[posicion + 1] if posicion + 1 < len(fila) else None
                resultado[campo] = _fecha(siguiente_crudo) or _fecha(siguiente)
            else:
                resultado[campo] = siguiente
    return resultado


def _fila_encabezado(hoja):
    for fila, valores in enumerate(hoja.iter_rows(max_row=30, max_col=1, values_only=True), start=1):
        primera = _clave(valores[0])
        if 'fase de proyecto formativo' in primera:
            return fila
    raise ValueError('No se encontró la fila de encabezados de la planeación pedagógica.')


def _codigo_resultado(resultado):
    coincidencia = re.match(r'^\s*(\d{5,})\s*[-:]', _texto(resultado))
    return coincidencia.group(1) if coincidencia else ''


def normalizar_instructor(nombre_raw: str, competencia: str = '', tipo_competencia: str = '') -> str:
    """Normaliza y unifica el nombre o perfil del instructor asignado en la planeación."""
    if not nombre_raw:
        if tipo_competencia == 'tecnica':
            return 'Instructor Técnico'
        if tipo_competencia == 'bilinguismo':
            return 'Instructor de Bilingüismo (Inglés)'
        if tipo_competencia == 'transversal':
            return 'Instructor Transversal'
        return 'Sin instructor asignado'

    texto = _texto(nombre_raw).strip()
    # Eliminar prefijos de conteo/viñeta tipo "1 ", "2 ", "1. ", "1- ", "1.- ", "1 / "
    texto = re.sub(r'^\s*\d+[\s\.\-_/:]+\s*', '', texto).strip()
    texto = re.sub(r'\s+', ' ', texto)
    if not texto:
        return 'Instructor Técnico' if tipo_competencia == 'tecnica' else 'Sin instructor asignado'

    clave = _clave(texto)

    # Perfiles genéricos conocidos del SENA
    if clave in ('instructor', 'instructor tecnico', 'instructores tecnicos', 'docente tecnico', 'tecnico'):
        return 'Instructor Técnico'

    if any(k in clave for k in ('ingles', 'english', 'bilinguismo', 'bilingue', 'lengua extranjera')):
        return 'Instructor de Bilingüismo (Inglés)'

    if any(k in clave for k in ('comunicacion', 'expresion', 'lenguaje', 'redaccion')):
        return 'Instructor de Comunicación'

    if any(k in clave for k in ('matematica', 'matematicas', 'cuantitativo', 'razonamiento')):
        return 'Instructor de Matemáticas'

    if any(k in clave for k in ('actividad fisica', 'deporte', 'habitos saludables', 'cultura fisica', 'educacion fisica')):
        return 'Instructor de Actividad Física'

    if any(k in clave for k in ('derechos fundamentales', 'derecho laboral', 'trabajo digno')):
        return 'Instructor de Derechos Fundamentales del Trabajo'

    if any(k in clave for k in ('salud ocupacional', 'seguridad y salud', 'sst', 'seguridad en el trabajo', 'salud en el trabajo')):
        return 'Instructor de SST / Salud Ocupacional'

    if any(k in clave for k in ('bienestar', 'induccion', 'orientador', 'psicologia', 'trabajo social')):
        return 'Bienestar al Aprendiz / Inducción'

    if any(k in clave for k in ('tic', 'informatica', 'herramientas informaticas', 'ofimatica', 'sistemas')):
        return 'Instructor de TIC'

    if any(k in clave for k in ('etica', 'valores', 'ciudadania', 'cultura de paz', 'enrique low')):
        return 'Instructor de Ética y Ciudadanía'

    if any(k in clave for k in ('ambiental', 'medio ambiente', 'ecologia')):
        return 'Instructor Ambiental'

    if any(k in clave for k in ('emprendimiento', 'innovacion', 'empresarismo', 'creacion de empresas')):
        return 'Instructor de Emprendimiento'

    # Si es nombre propio de una persona
    palabras = texto.split()
    return ' '.join(p.capitalize() for p in palabras)


def componente_instructor(nombre_normalizado: str) -> dict:
    """Devuelve el componente pedagógico, tono visual y badge para un instructor."""
    nom_l = (nombre_normalizado or '').lower()
    if 'técnico' in nom_l or 'tecnico' in nom_l:
        return {'tipo': 'tecnico', 'badge': 'Técnico', 'tono': 'primary', 'icono': '💻'}
    if 'inglés' in nom_l or 'ingles' in nom_l or 'bilingüismo' in nom_l:
        return {'tipo': 'bilinguismo', 'badge': 'Bilingüismo', 'tono': 'info', 'icono': '🌐'}
    if 'bienestar' in nom_l or 'inducción' in nom_l or 'induccion' in nom_l:
        return {'tipo': 'institucional', 'badge': 'Institucional', 'tono': 'neutral', 'icono': '🌱'}
    return {'tipo': 'transversal', 'badge': 'Transversal', 'tono': 'warning', 'icono': '📚'}


def parsear_planeacion(ruta):
    """Lee la hoja PLANEACION y agrupa filas repetidas del mismo resultado."""
    libro = openpyxl.load_workbook(Path(ruta), read_only=True, data_only=True)
    try:
        nombre_hoja = 'PLANEACION' if 'PLANEACION' in libro.sheetnames else libro.sheetnames[0]
        hoja = libro[nombre_hoja]
        mapa = _mapa_combinadas(hoja)
        fila_encabezado = _fila_encabezado(hoja)
        unidades_por_clave = {}
        filas_validas = 0
        total_directas = total_independientes = 0.0

        for fila, valores in _filas_combinadas(hoja, mapa):
            if fila < fila_encabezado + 2:
                continue
            fase = _valor_combinado(valores, mapa, fila, 1)
            actividad = _valor_combinado(valores, mapa, fila, 2)
            competencia = _valor_combinado(valores, mapa, fila, 3)
            rap = _valor_combinado(valores, mapa, fila, 4)
            directas = _numero(valores[8])
            independientes = _numero(valores[9])
            if not rap or (not competencia and not actividad):
                continue
            if not (rap or directas or independientes):
                continue

            filas_validas += 1
            total_directas += directas
            total_independientes += independientes
            instructor = _valor_combinado(valores, mapa, fila, 15)
            trimestre = _valor_combinado(valores, mapa, fila, 17)
            clave = (_clave(competencia), _clave(rap))
            unidad = unidades_por_clave.get(clave)
            if unidad is None:
                unidad = {
                    'fase': fase,
                    'actividad': actividad,
                    'competencia': competencia,
                    'competencia_tipo': clasificar_competencia(competencia),
                    'rap': rap,
                    'rap_codigo': _codigo_resultado(rap),
                    'horas_directas': 0,
                    'horas_independientes': 0,
                    'horas_total': 0,
                    'instructores': [],
                    'trimestres': [],
                    'filas_planificacion': 0,
                    'fila_inicial': fila,
                }
                unidades_por_clave[clave] = unidad
            unidad['horas_directas'] += directas
            unidad['horas_independientes'] += independientes
            unidad['horas_total'] += directas + independientes
            unidad['filas_planificacion'] += 1
            if instructor:
                inst_norm = normalizar_instructor(instructor, competencia, unidad['competencia_tipo'])
                if inst_norm and inst_norm not in unidad['instructores']:
                    unidad['instructores'].append(inst_norm)
            if trimestre and trimestre not in unidad['trimestres']:
                unidad['trimestres'].append(trimestre)

        unidades = sorted(unidades_por_clave.values(), key=lambda item: item['fila_inicial'])
        for unidad in unidades:
            unidad['horas_directas'] = round(unidad['horas_directas'], 2)
            unidad['horas_independientes'] = round(unidad['horas_independientes'], 2)
            unidad['horas_total'] = round(unidad['horas_total'], 2)
            unidad['trimestre'] = ', '.join(unidad['trimestres'])
            unidad.pop('fila_inicial', None)

        metadata = _metadatos(hoja)
        return {
            'metadata': metadata,
            'unidades': unidades,
            'resumen': {
                'filas_planificacion': filas_validas,
                'resultados': len(unidades),
                'competencias': len({item['competencia'] for item in unidades if item['competencia']}),
                'actividades': len({item['actividad'] for item in unidades if item['actividad']}),
                'fases': len({item['fase'] for item in unidades if item['fase']}),
                'horas_directas': round(total_directas, 2),
                'horas_independientes': round(total_independientes, 2),
                'horas_total': round(total_directas + total_independientes, 2),
            },
        }
    finally:
        libro.close()


def calcular_fechas_estimadas(unidades, fecha_inicio, fecha_fin):
    """Distribuye las fechas por horas de trabajo, con una aproximación auditable."""
    resultado = deepcopy(unidades)
    if not fecha_inicio or not fecha_fin or fecha_fin < fecha_inicio or not resultado:
        return resultado
    dias_totales = (fecha_fin - fecha_inicio).days + 1
    pesos = [max(_numero(item.get('horas_total')), 0) for item in resultado]
    total_horas = sum(pesos) or len(resultado)
    acumulado = 0.0
    for indice, item in enumerate(resultado):
        peso = pesos[indice] or 1
        inicio_offset = math.floor(acumulado / total_horas * dias_totales)
        acumulado += peso
        fin_offset = math.floor(acumulado / total_horas * dias_totales) - 1
        inicio_offset = min(max(inicio_offset, 0), dias_totales - 1)
        fin_offset = min(max(fin_offset, inicio_offset), dias_totales - 1)
        item['fecha_inicio_estimada'] = fecha_inicio + timedelta(days=inicio_offset)
        item['fecha_fin_estimada'] = fecha_inicio + timedelta(days=fin_offset)
    return resultado


def _codigo_normalizado(valor):
    coincidencia = re.search(r'\d{4,}', _texto(valor))
    return coincidencia.group(0) if coincidencia else ''


def _nombre_similitud(izquierda, derecha):
    a = _clave(izquierda)
    b = _clave(derecha)
    if not a or not b:
        return 1.0
    if a in b or b in a:
        return 1.0
    return SequenceMatcher(None, a, b).ratio()


def comparar_fuentes(metadata_planeacion, metadata_ficha):
    """Indica si la planeación y la ficha parecen pertenecer al mismo programa."""
    planeacion = metadata_planeacion or {}
    ficha = metadata_ficha or {}
    motivos = []
    codigo_plan = _codigo_normalizado(planeacion.get('codigo_programa'))
    codigo_ficha = _codigo_normalizado(ficha.get('codigo_programa'))
    nombre_plan = planeacion.get('nombre_programa') or ''
    nombre_ficha = ficha.get('nombre_programa') or ''

    if codigo_plan and codigo_ficha and codigo_plan != codigo_ficha:
        motivos.append(
            f'El código del programa no coincide: el archivo cargado corresponde al programa {codigo_plan}'
            + (f' ("{nombre_plan}")' if nombre_plan else '')
            + f' y la ficha actual está configurada para el programa {codigo_ficha}'
            + (f' ("{nombre_ficha}")' if nombre_ficha else '')
            + '.'
        )
    if nombre_plan and nombre_ficha and _nombre_similitud(nombre_plan, nombre_ficha) < 0.55:
        motivos.append(
            f'El nombre del programa no coincide: planeación "{nombre_plan}" vs ficha "{nombre_ficha}".'
        )
    return {
        'alineado': not motivos,
        'motivos': motivos,
        'codigo_planeacion': codigo_plan,
        'codigo_ficha': codigo_ficha,
        'nombre_planeacion': nombre_plan,
        'nombre_ficha': nombre_ficha,
    }

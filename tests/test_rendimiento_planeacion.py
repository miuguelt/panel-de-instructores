"""Regresiones de los recorridos costosos encontrados en producción."""

import io
import hashlib
import unittest
from datetime import date
from difflib import SequenceMatcher
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from werkzeug.datastructures import FileStorage

from app.services import emparejamiento_juicios as cruce
from app.services import importacion_ficha as importacion
from app import db
from app.models import ArchivoFichaVersion
from app.services import fases_dashboard, versiones_archivos
from app.tyt import calculo as tyt
from tests.apoyo_planeacion import crear_planeacion_xlsx
from tests.test_rendimiento import BaseRendimiento


class CorrespondenciaRapTestCase(unittest.TestCase):
    def test_no_crea_comparadores_para_textos_descartados_por_las_cotas(self):
        registro = object()
        with patch.object(cruce, 'SequenceMatcher', wraps=SequenceMatcher) as comparadores:
            self.assertEqual(cruce.juicios_de({'rap': 'abcdefghij'},
                {'zz': [], 'zzzzzzzzzz': [], 'abcdefghix': [registro]}, {}), [registro])
        self.assertEqual(comparadores.call_count, 1)

    def test_normalizacion_reutiliza_textos_sin_cambiar_reglas(self):
        cruce._normalizar_clave.cache_clear()
        normalizar = cruce.unicodedata.normalize
        with patch.object(cruce.unicodedata, 'normalize', wraps=normalizar) as llamadas:
            for _ in range(20):
                self.assertEqual(cruce.clave_resultado('123456 - Ámbito del caché de RAP.'), 'ambito del cache de rap')
        self.assertEqual(llamadas.call_count, 1)
        self.assertEqual(cruce.clave_resultado(['Á', 'B']), 'a b')

    def test_tyt_reutiliza_normalizacion_y_conserva_datos_mutables(self):
        tyt._normalizar_texto.cache_clear()
        normalizar = tyt.unicodedata.normalize
        with patch.object(tyt.unicodedata, 'normalize', wraps=normalizar) as llamadas:
            for _ in range(20):
                self.assertEqual(tyt._texto('  Ámbito del caché de TyT  '), 'AMBITO DEL CACHE DE TYT')
        self.assertEqual(llamadas.call_count, 1)
        valor = ['Á']
        self.assertEqual(tyt._texto(valor), "['A']")
        valor.append('B')
        self.assertEqual(tyt._texto(valor), "['A', 'B']")

    def test_codigo_y_texto_exacto_conservan_prioridad_y_registros(self):
        por_codigo = SimpleNamespace(resultado_aprendizaje='123456 - Texto del reporte')
        por_texto = SimpleNamespace(resultado_aprendizaje='Téxto   de la planeación.')
        claves, codigos = cruce.indexar_juicios([por_codigo, por_texto])
        with patch.object(SequenceMatcher, 'ratio', autospec=True,
                          side_effect=SequenceMatcher.ratio) as costoso:
            self.assertEqual(cruce.juicios_de({'rap_codigo': '123456', 'rap': 'Otro'}, claves, codigos), [por_codigo])
            self.assertEqual(cruce.juicios_de({'rap': 'texto de la planeacion'}, claves, codigos), [por_texto])
        costoso.assert_not_called()
        self.assertEqual(cruce.texto_limpio(' -- '), '')
        self.assertEqual(cruce.codigo_resultado('Sin código'), '')
        self.assertTrue(cruce.es_aprobado('Aprobada'))
        self.assertFalse(cruce.es_aprobado('NO APROBADO'))

    def test_descarta_imposibles_sin_calcular_alineamientos(self):
        registro = object()
        claves = {'zz': [object()], 'zzzzzzzzzz': [object()], 'abcdefghix': [registro]}
        with patch.object(SequenceMatcher, 'ratio', autospec=True,
                          side_effect=SequenceMatcher.ratio) as costoso:
            resultado = cruce.juicios_de({'rap': 'abcdefghij'}, claves, {})
        self.assertEqual(resultado, [registro], 'El candidato en el umbral del 90 % se conserva.')
        self.assertEqual(costoso.call_count, 1, 'Solo el candidato viable necesita alineamiento.')

    def test_ambiguedad_detiene_la_busqueda_al_segundo_candidato(self):
        claves = {texto: [object()] for texto in ('abcdefghix', 'abcdefghiy', 'abcdefghiz')}
        with patch.object(SequenceMatcher, 'ratio', autospec=True,
                          side_effect=SequenceMatcher.ratio) as costoso:
            self.assertEqual(cruce.juicios_de({'rap': 'abcdefghij'}, claves, {}), [])
        self.assertEqual(costoso.call_count, 2)

    def test_resultados_equivalen_a_la_regla_original_en_limites_y_textos_largos(self):
        textos = ['abcdefghix', 'abcdefgxyz', '', 'interpretar requisitos del negocio ' * 9,
                  'interpretar requerimientos del negocio ' * 9, 'sistemas y seguridad']
        claves = {texto: [object()] for texto in textos}
        for rap in textos + ['abcdefghij', 'ausente', None, 'Interpretar requisitos del negocio ' * 9]:
            with self.subTest(rap=rap):
                clave = cruce.clave_resultado(rap)
                esperados = claves.get(clave, [])
                if not esperados:
                    candidatos = [registros for otra, registros in claves.items()
                                  if SequenceMatcher(None, clave, otra).ratio() >= cruce.SIMILITUD_MINIMA]
                    esperados = candidatos[0] if len(candidatos) == 1 else []
                self.assertEqual(cruce.juicios_de({'rap': rap}, claves, {}), esperados)
        self.assertEqual(cruce.juicios_de({'rap': None}, {}, {}), [])


class MetadataAcotadaTestCase(unittest.TestCase):
    def test_xls_real_acota_metadata_y_conserva_importacion_completa(self):
        import xlrd
        contenido = (Path(__file__).parent / 'fixtures/rendimiento_metadatos.xls').read_bytes()
        archivo = FileStorage(stream=io.BytesIO(contenido), filename='reporte.xls')
        leer_fila = xlrd.sheet.Sheet.row_values
        with patch.object(xlrd.sheet.Sheet, 'row_values', autospec=True,
                          side_effect=leer_fila) as filas:
            metadata = importacion.leer_metadata_archivo(archivo)
        self.assertEqual(metadata['codigo_ficha'], '3999999')
        self.assertEqual(metadata['fecha_reporte'].isoformat(), '2026-10-01')
        self.assertEqual(filas.call_count, 20)
        self.assertEqual(archivo.stream.tell(), 0)
        metadata_completa, registros = importacion._leer_archivo(archivo)
        self.assertEqual(metadata_completa, metadata)
        self.assertEqual(len(registros), 35)
        self.assertEqual(registros[-1]['documento'], '34')

    def test_xls_cierra_recursos_al_fallar_metadata(self):
        import xlrd
        contenido = (Path(__file__).parent / 'fixtures/rendimiento_metadatos.xls').read_bytes()
        archivo = FileStorage(stream=io.BytesIO(contenido), filename='reporte.xls')
        cerrar = xlrd.book.Book.release_resources
        with patch.object(xlrd.book.Book, 'release_resources', autospec=True,
                          side_effect=cerrar) as recursos, \
                patch.object(importacion, '_extraer_metadata', side_effect=ValueError('Error de lectura')):
            with self.assertRaisesRegex(ValueError, 'Error de lectura'):
                importacion.leer_metadata_archivo(archivo)
        recursos.assert_called_once()
        self.assertEqual(archivo.stream.tell(), 0)

    def crear_archivo(self, cantidad=2000):
        libro = openpyxl.Workbook()
        hoja = libro.active
        hoja.append(['Ficha de Caracterización:', None, '3999999'])
        hoja.append(['Código:', '228118'])
        hoja.append(['Fecha del Reporte:', '01/10/2026'])
        hoja.append(['Denominación:', 'Programa de prueba'])
        hoja.append(['Documento', 'Nombre', 'Resultado', 'Juicio'])
        for numero in range(cantidad):
            hoja.append([str(numero), 'Aprendiz de prueba', 'RAP de prueba', 'APROBADO'])
        stream = io.BytesIO()
        libro.save(stream)
        libro.close()
        stream.seek(0)
        return FileStorage(stream=stream, filename='reporte.xlsx')

    def test_metadatos_lee_veinte_filas_y_no_construye_registros(self):
        archivo = self.crear_archivo()
        filas_leidas = []
        iterar = ReadOnlyWorksheet.iter_rows

        def registrar(hoja, *args, **kwargs):
            for fila in iterar(hoja, *args, **kwargs):
                filas_leidas.append(fila)
                yield fila

        with patch.object(ReadOnlyWorksheet, 'iter_rows', registrar), \
                patch.object(importacion, '_parse_rows', wraps=importacion._parse_rows) as registros:
            metadata = importacion.leer_metadata_archivo(archivo)
        self.assertEqual(metadata['codigo_ficha'], '3999999')
        self.assertEqual(metadata['codigo_programa'], '228118')
        self.assertEqual(metadata['nombre_programa'], 'Programa de prueba')
        self.assertEqual(metadata['fecha_reporte'].isoformat(), '2026-10-01')
        self.assertEqual(len(filas_leidas), 20)
        registros.assert_not_called()
        self.assertEqual(archivo.stream.tell(), 0)

    def test_importacion_completa_conserva_todas_las_filas(self):
        metadata, registros = importacion._leer_archivo(self.crear_archivo(cantidad=35))
        self.assertEqual(metadata['codigo_ficha'], '3999999')
        self.assertEqual(len(registros), 35)
        self.assertEqual(registros[-1]['documento'], '34')

    def test_error_cierra_libro_y_restaura_posicion_del_stream(self):
        archivo = self.crear_archivo(cantidad=1)
        cargar = openpyxl.load_workbook
        libros = []

        def registrar(*args, **kwargs):
            libro = cargar(*args, **kwargs)
            libros.append(libro)
            return libro

        with patch.object(openpyxl, 'load_workbook', side_effect=registrar), \
                patch.object(importacion, '_extraer_metadata', side_effect=ValueError('Error de lectura')):
            with self.assertRaisesRegex(ValueError, 'Error de lectura'):
                importacion.leer_metadata_archivo(archivo)
        self.assertEqual(archivo.stream.tell(), 0)
        self.assertIsNone(libros[0]._archive.fp, 'El libro se cierra también al fallar el lector.')

    def test_excel_invalido_conserva_error_y_posicion(self):
        for extension, mensaje in [('xlsx', 'Excel válido'), ('xls', '.xls no es válido')]:
            with self.subTest(extension=extension):
                archivo = FileStorage(stream=io.BytesIO(b'Excel invalido'), filename=f'reporte.{extension}')
                archivo.stream.seek(3)
                with self.assertRaisesRegex(importacion.ErrorImportacion, mensaje):
                    importacion.leer_metadata_archivo(archivo)
                self.assertEqual(archivo.stream.tell(), 3)


class LecturaFasesTestCase(BaseRendimiento):
    def test_consultar_ultima_version_sin_recuperacion_no_genera_archivos(self):
        with patch.object(versiones_archivos, 'asegurar_version_reporte',
                          wraps=versiones_archivos.asegurar_version_reporte) as recuperar:
            version = versiones_archivos.ultima_version(
                self.ficha.id, 'reporte_juicios', solo_procesadas=True, recuperar_reporte=False)
        self.assertIsNone(version)
        recuperar.assert_not_called()
        self.assertEqual(ArchivoFichaVersion.query.count(), 0)
        self.assertEqual(list(Path(self.uploads.name).rglob('*.xlsx')), [])

    def test_fases_usan_version_registrada_sin_buscar_reportes_en_disco(self):
        ruta = crear_planeacion_xlsx(self.uploads.name)
        self.ficha.fecha_fin = date(2026, 12, 31)
        db.session.add(ArchivoFichaVersion(
            ficha_id=self.ficha.id, instructor_id=self.ficha.instructor_id,
            tipo='planeacion', version=1, estado='procesado',
            nombre_archivo=ruta.name, ruta_archivo=ruta.name, tamano_bytes=ruta.stat().st_size,
            hash_sha256=hashlib.sha256(ruta.read_bytes()).hexdigest()))
        db.session.commit()
        fases_dashboard._CACHE_PLANEACION_PARSED.clear()
        with patch.object(versiones_archivos, 'asegurar_version_reporte',
                          wraps=versiones_archivos.asegurar_version_reporte) as recuperar:
            resultado = fases_dashboard.obtener_seguimiento_fases_dashboard(self.ficha, hoy=date(2026, 1, 10))
        self.assertTrue(resultado['disponible'])
        self.assertEqual(resultado['fuente'], 'planeacion_gfpi')
        self.assertEqual(resultado['resumen_raps']['total'], 2)
        self.assertEqual(resultado['resumen_raps']['aprobados'], 0)
        recuperar.assert_not_called()
        self.assertEqual(ArchivoFichaVersion.query.count(), 1)

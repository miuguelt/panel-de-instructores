"""Planeación pedagógica mínima en Excel para las pruebas del módulo."""

from datetime import date
from pathlib import Path

import openpyxl


def crear_planeacion_xlsx(carpeta):
    """Genera un GFPI-F-134 reducido con celdas combinadas y dos resultados."""
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = 'PLANEACION'
    hoja['A5'] = 'Fecha de Elaboración'
    hoja['B5'] = date(2025, 1, 21)
    hoja['A6'] = 'Denominación del Programa de Formación'
    hoja['B6'] = 'Tecnólogo en Analisis y Desarrollo de Software'
    hoja['A8'] = 'Código y versión del Programa de Formación'
    hoja['B8'] = '228118 v 1.0'
    hoja['A16'] = 'FASE DE PROYECTO FORMATIVO'
    hoja['B16'] = 'ACTIVIDAD DE PROYECTO FORMATIVO'
    hoja['C16'] = 'COMPETENCIA'
    hoja['D16'] = 'RESULTADOS DE APRENDIZAJE'
    hoja['I16'] = 'DURACIÓN ACTIVIDAD DE APRENDIZAJE (HORAS)'
    hoja['I17'] = 'HORAS TRABAJO DIRECTO'
    hoja['J17'] = 'HORAS TRABAJO INDEPENDIENTE'
    hoja['O16'] = 'INSTRUCTORES RESPONSABLES'
    hoja['Q16'] = 'TRIMESTRE'

    hoja['A18'] = 'ANÁLISIS'
    hoja['B18'] = 'AP01. Contextualizar'
    hoja['C18'] = 'Inducción'
    hoja['D18'] = '601390 - Identificar la dinámica organizacional del SENA.'
    hoja['I18'] = 10
    hoja['J18'] = 2
    hoja['O18'] = 'Bienestar al aprendiz'
    hoja['Q18'] = 'Trimestre 1'

    hoja['D19'] = '601390 - Identificar la dinámica organizacional del SENA.'
    hoja['I19'] = 8
    hoja['J19'] = 2
    hoja['O19'] = 'Instructor técnico'

    hoja['A20'] = 'PLANEACIÓN'
    hoja['B20'] = 'AP02. Especificar'
    hoja['C20'] = 'Comunicación oral y escrita'
    hoja['D20'] = '601401 - Comunicar de manera eficaz.'
    hoja['I20'] = 20
    hoja['J20'] = 5
    hoja['O20'] = 'Instructor Comunicación'
    hoja['Q20'] = 'Trimestre 2'

    hoja.merge_cells('A18:A19')
    hoja.merge_cells('B18:B19')
    hoja.merge_cells('C18:C19')
    hoja.merge_cells('Q18:Q19')
    ruta = Path(carpeta) / 'planeacion.xlsx'
    libro.save(ruta)
    return ruta

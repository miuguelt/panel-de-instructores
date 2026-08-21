import unittest
from datetime import date
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

import openpyxl

from app import create_app, db
from app.models import ArchivoFichaVersion, Ficha, Instructor
from tests.apoyo_planeacion import crear_planeacion_xlsx
from tests.apoyo_programa import crear_programa_pdf


class PlaneacionRouteTestCase(unittest.TestCase):
    def setUp(self):
        self.uploads = TemporaryDirectory()
        self.app = create_app({
            'TESTING': True,
            'SQLALCHEMY_DATABASE_URI': 'sqlite:///:memory:',
            'SQLALCHEMY_ENGINE_OPTIONS': {},
            'WTF_CSRF_ENABLED': False,
            'RATELIMIT_ENABLED': False,
            'UPLOAD_FOLDER': self.uploads.name,
        })
        self.contexto = self.app.app_context()
        self.contexto.push()
        db.create_all()
        self.instructor = Instructor(nombre='Uno', correo='uno-planeacion@sena.edu.co')
        self.instructor.set_password('x')
        db.session.add(self.instructor)
        db.session.flush()
        self.ficha = Ficha(
            codigo='3336360',
            codigo_ficha='3336360',
            codigo_programa='228118',
            nombre_programa='Tecnólogo en Analisis y Desarrollo de Software',
            instructor_id=self.instructor.id,
            fecha_inicio=date(2026, 1, 1),
            fecha_fin=date(2026, 12, 31),
        )
        db.session.add(self.ficha)
        db.session.commit()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.contexto.pop()
        self.uploads.cleanup()

    def test_carga_planeacion_crea_version_y_habilita_descarga(self):
        ruta = crear_planeacion_xlsx(self.uploads.name)
        cliente = self.app.test_client()
        cliente.post('/login', data={'correo': self.instructor.correo, 'password': 'x'})
        with ruta.open('rb') as stream:
            pagina = cliente.post(
                f'/instructor/fichas/{self.ficha.id}/planeacion/cargar',
                data={'archivo_planeacion': (stream, ruta.name)},
                content_type='multipart/form-data',
                follow_redirects=True,
            )
        version = ArchivoFichaVersion.query.one()
        descarga = cliente.get(
            f'/instructor/fichas/{self.ficha.id}/planeacion/archivos/{version.id}/descargar'
        )
        cuerpo = pagina.get_data(as_text=True)
        self.assertEqual(pagina.status_code, 200)
        self.assertIn('Línea de tiempo de la ficha', cuerpo)
        self.assertIn('Cronograma de la ficha: plan vs. avance real', cuerpo)
        self.assertIn('Etapa productiva', cuerpo)
        self.assertIn('Detalle por fase del proyecto formativo', cuerpo)
        self.assertEqual(version.estado, 'procesado')
        self.assertEqual(descarga.status_code, 200)
        self.assertIn(b'attachment', descarga.headers.get('Content-Disposition', '').encode())
        descarga.close()

    def test_no_duplica_una_planeacion_con_el_mismo_contenido(self):
        ruta = crear_planeacion_xlsx(self.uploads.name)
        cliente = self.app.test_client()
        cliente.post('/login', data={'correo': self.instructor.correo, 'password': 'x'})
        with ruta.open('rb') as stream:
            cliente.post(
                f'/instructor/fichas/{self.ficha.id}/planeacion/cargar',
                data={'archivo_planeacion': (stream, ruta.name)},
                content_type='multipart/form-data',
                follow_redirects=True,
            )
        with ruta.open('rb') as stream:
            pagina = cliente.post(
                f'/instructor/fichas/{self.ficha.id}/planeacion/cargar',
                data={'archivo_planeacion': (stream, ruta.name)},
                content_type='multipart/form-data',
                follow_redirects=True,
            )

        self.assertEqual(ArchivoFichaVersion.query.count(), 1)
        self.assertIn('ya fue cargado', pagina.get_data(as_text=True))

    def test_carga_con_xhr_devuelve_resultado_json_sin_redireccion_intermedia(self):
        ruta = crear_planeacion_xlsx(self.uploads.name)
        cliente = self.app.test_client()
        cliente.post('/login', data={'correo': self.instructor.correo, 'password': 'x'})
        with ruta.open('rb') as stream:
            respuesta = cliente.post(
                f'/instructor/fichas/{self.ficha.id}/planeacion/cargar',
                data={'archivo_planeacion': (stream, ruta.name)},
                content_type='multipart/form-data',
                headers={'X-Requested-With': 'XMLHttpRequest'},
            )

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta.json['ok'], True)
        self.assertIn('/planeacion', respuesta.json['redirect'])
        self.assertEqual(ArchivoFichaVersion.query.count(), 1)

    def test_carga_con_xhr_informa_en_json_una_planeacion_desalineada(self):
        ruta = crear_planeacion_xlsx(self.uploads.name)
        libro = openpyxl.load_workbook(ruta)
        libro['PLANEACION']['B8'] = '999999 v 1.0'
        desalineada = Path(self.uploads.name) / 'planeacion_otro_programa_xhr.xlsx'
        libro.save(desalineada)

        cliente = self.app.test_client()
        cliente.post('/login', data={'correo': self.instructor.correo, 'password': 'x'})
        with desalineada.open('rb') as stream:
            respuesta = cliente.post(
                f'/instructor/fichas/{self.ficha.id}/planeacion/cargar',
                data={'archivo_planeacion': (stream, desalineada.name)},
                content_type='multipart/form-data',
                headers={'X-Requested-With': 'XMLHttpRequest'},
            )

        self.assertEqual(respuesta.status_code, 422)
        self.assertEqual(respuesta.json['ok'], False)
        self.assertIn('Carga bloqueada', respuesta.json['message'])
        self.assertEqual(ArchivoFichaVersion.query.count(), 0)

    def test_bloquea_planeacion_de_otro_programa(self):
        ruta = crear_planeacion_xlsx(self.uploads.name)
        libro = openpyxl.load_workbook(ruta)
        libro['PLANEACION']['B8'] = '999999 v 1.0'
        desalineada = Path(self.uploads.name) / 'planeacion_otro_programa.xlsx'
        libro.save(desalineada)

        cliente = self.app.test_client()
        cliente.post('/login', data={'correo': self.instructor.correo, 'password': 'x'})
        with desalineada.open('rb') as stream:
            pagina = cliente.post(
                f'/instructor/fichas/{self.ficha.id}/planeacion/cargar',
                data={'archivo_planeacion': (stream, desalineada.name)},
                content_type='multipart/form-data',
                follow_redirects=True,
            )

        self.assertEqual(ArchivoFichaVersion.query.count(), 0)
        self.assertIn('Carga bloqueada', pagina.get_data(as_text=True))

    def test_bloquea_planeacion_si_la_ficha_no_tiene_codigo_de_programa(self):
        self.ficha.codigo_programa = None
        db.session.commit()
        ruta = crear_planeacion_xlsx(self.uploads.name)

        cliente = self.app.test_client()
        cliente.post('/login', data={'correo': self.instructor.correo, 'password': 'x'})
        with ruta.open('rb') as stream:
            pagina = cliente.post(
                f'/instructor/fichas/{self.ficha.id}/planeacion/cargar',
                data={'archivo_planeacion': (stream, ruta.name)},
                content_type='multipart/form-data',
                follow_redirects=True,
            )

        self.assertEqual(ArchivoFichaVersion.query.count(), 0)
        self.assertIn('código del programa', pagina.get_data(as_text=True))

    def test_bloquea_reporte_de_otra_ficha_antes_de_importar(self):
        libro = openpyxl.Workbook()
        hoja = libro.active
        hoja.append(['Ficha de Caracterización:', '', '9999999'])
        hoja.append(['Código:', '', '228118'])
        hoja.append(['Denominación:', '', 'Tecnólogo en Analisis y Desarrollo de Software'])
        for _ in range(10):
            hoja.append([])
        hoja.append([
            'Tipo de Documento', 'Número de Documento', 'Nombre', 'Apellidos', 'Estado',
            'Competencia', 'Resultado de Aprendizaje', 'Juicio de Evaluación',
            'Fecha y Hora del Juicio Evaluativo', 'Funcionario que registro el juicio evaluativo',
        ])
        hoja.append([
            'CC', '1001', 'Ana', 'Pérez', 'EN_FORMACION', 'Competencia 1', 'Resultado 1',
            'APROBADO', '14/03/2023', 'Instructor anterior',
        ])
        salida = BytesIO()
        libro.save(salida)
        salida.seek(0)

        cliente = self.app.test_client()
        cliente.post('/login', data={'correo': self.instructor.correo, 'password': 'x'})
        pagina = cliente.post(
            f'/instructor/fichas/{self.ficha.id}/cargar-excel',
            data={'archivo': (salida, 'Reporte de Juicios Evaluativos.xlsx'), 'retorno': 'planeacion'},
            content_type='multipart/form-data',
            follow_redirects=True,
        )

        self.assertEqual(ArchivoFichaVersion.query.count(), 0)
        self.assertIn('Carga bloqueada', pagina.get_data(as_text=True))

    def test_carga_el_programa_de_formacion_y_lo_deja_descargable(self):
        ruta = crear_programa_pdf(self.uploads.name)
        cliente = self.app.test_client()
        cliente.post('/login', data={'correo': self.instructor.correo, 'password': 'x'})
        with ruta.open('rb') as stream:
            pagina = cliente.post(
                f'/instructor/fichas/{self.ficha.id}/planeacion/programa',
                data={'archivo_programa': (stream, ruta.name)},
                content_type='multipart/form-data',
                follow_redirects=True,
            )

        version = ArchivoFichaVersion.query.filter_by(tipo='programa_formacion').one()
        descarga = cliente.get(
            f'/instructor/fichas/{self.ficha.id}/planeacion/archivos/{version.id}/descargar'
        )
        self.assertEqual(pagina.status_code, 200)
        self.assertEqual(version.estado, 'procesado')
        self.assertIn('Programa de formación', pagina.get_data(as_text=True))
        self.assertEqual(descarga.status_code, 200)
        descarga.close()

    def test_rechaza_un_programa_que_no_es_pdf(self):
        ruta = crear_planeacion_xlsx(self.uploads.name)
        cliente = self.app.test_client()
        cliente.post('/login', data={'correo': self.instructor.correo, 'password': 'x'})
        with ruta.open('rb') as stream:
            pagina = cliente.post(
                f'/instructor/fichas/{self.ficha.id}/planeacion/programa',
                data={'archivo_programa': (stream, ruta.name)},
                content_type='multipart/form-data',
                follow_redirects=True,
            )

        self.assertEqual(ArchivoFichaVersion.query.filter_by(tipo='programa_formacion').count(), 0)
        self.assertIn('formato .pdf', pagina.get_data(as_text=True))

    def test_carga_todo_los_tres_documentos_en_una_sola_accion(self):
        ruta_plan = crear_planeacion_xlsx(self.uploads.name)
        ruta_prog = crear_programa_pdf(self.uploads.name)

        libro = openpyxl.Workbook()
        hoja = libro.active
        hoja.append(['Ficha de Caracterización:', '', self.ficha.codigo])
        hoja.append(['Código:', '', '228118'])
        hoja.append(['Denominación:', '', 'Tecnólogo en Analisis y Desarrollo de Software'])
        for _ in range(10):
            hoja.append([])
        hoja.append([
            'Tipo de Documento', 'Número de Documento', 'Nombre', 'Apellidos', 'Estado',
            'Competencia', 'Resultado de Aprendizaje', 'Juicio de Evaluación',
            'Fecha y Hora del Juicio Evaluativo', 'Funcionario que registro el juicio evaluativo',
        ])
        hoja.append([
            'CC', '1001', 'Ana', 'Pérez', 'EN_FORMACION', 'Competencia 1', 'Resultado 1',
            'APROBADO', '14/03/2023', 'Instructor anterior',
        ])
        salida_juicios = BytesIO()
        libro.save(salida_juicios)
        salida_juicios.seek(0)

        cliente = self.app.test_client()
        cliente.post('/login', data={'correo': self.instructor.correo, 'password': 'x'})

        with ruta_plan.open('rb') as s_plan, ruta_prog.open('rb') as s_prog:
            pagina = cliente.post(
                f'/instructor/fichas/{self.ficha.id}/planeacion/cargar-todo',
                data={
                    'archivo_planeacion': (s_plan, ruta_plan.name),
                    'archivo_programa': (s_prog, ruta_prog.name),
                    'archivo_juicios': (salida_juicios, 'Reporte.xlsx'),
                },
                content_type='multipart/form-data',
                follow_redirects=True,
            )

        self.assertEqual(pagina.status_code, 200)
        self.assertEqual(ArchivoFichaVersion.query.count(), 3)
        cuerpo = pagina.get_data(as_text=True)
        self.assertIn('Línea de tiempo de la ficha', cuerpo)
        # Con las tres fuentes sincronizadas, el centro de carga se muestra compacto
        self.assertIn('Fuentes sincronizadas', cuerpo)
        self.assertIn('Actualizar archivos fuente', cuerpo)

    def test_carga_todo_xhr_responde_json(self):
        ruta_plan = crear_planeacion_xlsx(self.uploads.name)
        cliente = self.app.test_client()
        cliente.post('/login', data={'correo': self.instructor.correo, 'password': 'x'})

        with ruta_plan.open('rb') as s_plan:
            respuesta = cliente.post(
                f'/instructor/fichas/{self.ficha.id}/planeacion/cargar-todo',
                data={'archivo_planeacion': (s_plan, ruta_plan.name)},
                content_type='multipart/form-data',
                headers={'X-Requested-With': 'XMLHttpRequest'},
            )

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta.json['ok'], True)
        self.assertIn('/planeacion', respuesta.json['redirect'])

    def test_descargar_todos_fuentes_en_zip(self):
        ruta_plan = crear_planeacion_xlsx(self.uploads.name)
        ruta_prog = crear_programa_pdf(self.uploads.name)

        cliente = self.app.test_client()
        cliente.post('/login', data={'correo': self.instructor.correo, 'password': 'x'})

        # Sin archivos -> redirección con advertencia
        sin_archivos = cliente.get(f'/instructor/fichas/{self.ficha.id}/planeacion/descargar-todos', follow_redirects=True)
        self.assertIn('No hay archivos cargados', sin_archivos.get_data(as_text=True))

        # Cargamos archivos
        with ruta_plan.open('rb') as s_plan, ruta_prog.open('rb') as s_prog:
            cliente.post(
                f'/instructor/fichas/{self.ficha.id}/planeacion/cargar-todo',
                data={
                    'archivo_planeacion': (s_plan, ruta_plan.name),
                    'archivo_programa': (s_prog, ruta_prog.name),
                },
                content_type='multipart/form-data',
                follow_redirects=True,
            )

        resp_zip = cliente.get(f'/instructor/fichas/{self.ficha.id}/planeacion/descargar-todos')
        self.assertEqual(resp_zip.status_code, 200)
        self.assertEqual(resp_zip.mimetype, 'application/zip')
        self.assertIn(f'fuentes_ficha_{self.ficha.codigo}.zip', resp_zip.headers.get('Content-Disposition', ''))


if __name__ == '__main__':
    unittest.main()



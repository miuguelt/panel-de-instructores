"""Regresiones verificadas con registros persistidos y documentos reales de prueba."""

from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError

import openpyxl
import pytest
from werkzeug.datastructures import FileStorage

from app import db
from app.models import Entrega, Ficha, Instructor, JuicioEvaluativo, Tarea
from app.models.aprendiz import normalizar_url_portafolio
from app.models.ranking import ConfiguracionRanking, PuntajeHistorico
from app.models.importacion import ImportacionJob
from app.models.grupo import Grupo
from app.services.atencion_service import AtencionService
from app.services.alertas import ejecutar_revision_automatica
from app.services.grupo_service import archivar_grupos_de_ficha, crear_grupos_aleatorios
from app.services.curva_rendimiento import obtener_curva_rendimiento_aprendiz
from app.services.importacion_ficha import _fecha, _leer_archivo
from app.services.importacion_jobs import ColaImportacionesNoDisponible
from app.services.panorama_planeacion import _calcular_desempeno_ficha
from app.services.ranking import calcular_ranking
from app.services.planeacion import parsear_planeacion
from app.services.recuperacion_reportes import asegurar_version_reporte
from app.services.recuperacion_planeacion import asegurar_version_planeacion, buscar_archivo_planeacion_original
from app.services.recuperacion_programa import asegurar_version_programa, buscar_archivo_programa_original
from app.services.recuperacion_fuentes import recuperar_fuentes_ficha
from app.services.tareas import agrupar_tareas_por_corte_e_instructor, agrupar_tareas_por_instructor
from tests import test_archivos as soporte
from tests.apoyo_planeacion import crear_planeacion_xlsx
from tests.apoyo_programa import crear_programa_pdf


@pytest.fixture
def contexto():
    caso = soporte.ArchivosTestCase()
    caso.setUp()
    try:
        yield caso
    finally:
        caso.tearDown()


@pytest.mark.parametrize('agrupador', ['instructor', 'corte'])
def test_instructores_homonimos_no_dividen_sus_tareas_en_varios_grupos(contexto, agrupador):
    contexto.instructor.nombre = 'Camila Pérez'
    homonimo = Instructor(nombre='Camila Pérez', correo='homonimo@example.com', password_hash='example')
    db.session.add(homonimo)
    db.session.flush()
    ahora = datetime(2026, 10, 21, 12)
    for indice, instructor in enumerate([contexto.instructor, homonimo, contexto.instructor]):
        db.session.add(Tarea(ficha_id=contexto.ficha.id, instructor_id=instructor.id,
                            titulo=f'Tarea {indice}', creada_en=ahora - timedelta(days=indice),
                            fecha_limite=ahora + timedelta(days=indice + 1)))
    db.session.commit()
    tareas = Tarea.query.order_by(Tarea.id).all()
    if agrupador == 'instructor':
        grupos = agrupar_tareas_por_instructor(tareas, ahora=ahora)
    else:
        grupos = agrupar_tareas_por_corte_e_instructor(tareas, [], ahora=ahora)[0]['instructores']
    assert len(grupos) == 2
    assert sorted(len(grupo['tareas']) for grupo in grupos) == [1, 2]
    assert {t.id for grupo in grupos for t in grupo['tareas']} == {t.id for t in tareas}


def test_reporte_reconstruido_conserva_identidad_y_fecha_academica(contexto):
    fecha_juicio = datetime(2025, 3, 14, 8, 35)
    db.session.add(JuicioEvaluativo(
        ficha_id=contexto.ficha.id, aprendiz_id=contexto.aprendiz.id,
        competencia='Construir software', resultado_aprendizaje='RAP 01', juicio='APROBADO',
        fecha_juicio=fecha_juicio, fecha_fuente_texto='14/03/2025 08:35',
        importado_en=datetime(2026, 10, 21, 12), funcionario_registro='Instructor de prueba',
    ))
    db.session.commit()
    version = asegurar_version_reporte(contexto.ficha)
    ruta = Path(contexto.uploads.name) / version.ruta_archivo
    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    filas = list(libro.active.values)
    libro.close()
    assert filas[1][8] == '14/03/2025 08:35:00'
    with ruta.open('rb') as archivo:
        _, registros = _leer_archivo(FileStorage(BytesIO(archivo.read()), filename=ruta.name))
    assert len(registros) == 1
    assert registros[0]['documento'] == contexto.aprendiz.documento
    assert registros[0]['fecha_juicio'] == fecha_juicio
    assert registros[0]['juicio'] == 'APROBADO'


def test_reporte_reconstruido_sin_fecha_no_inventa_la_fecha_de_importacion(contexto):
    db.session.add(JuicioEvaluativo(
        ficha_id=contexto.ficha.id, aprendiz_id=contexto.aprendiz.id,
        competencia='Construir software', resultado_aprendizaje='RAP 01', juicio='APROBADO',
        fecha_juicio=None, fecha_fuente_texto='Pendiente por confirmar',
        importado_en=datetime(2026, 10, 21, 12),
    ))
    db.session.commit()
    version = asegurar_version_reporte(contexto.ficha)
    libro = openpyxl.load_workbook(Path(contexto.uploads.name) / version.ruta_archivo, read_only=True)
    try:
        assert libro.active.cell(2, 9).value == 'Pendiente por confirmar'
    finally:
        libro.close()


@pytest.mark.parametrize('valor', [float('inf'), 10 ** 100, 1e12])
def test_fecha_excel_fuera_de_rango_no_interrumpe_la_importacion(valor):
    assert _fecha(valor) is None


@pytest.mark.parametrize('valor', ['NaN', 'inf', '-inf'])
def test_planeacion_rechaza_duraciones_no_finitas_en_excel_real(contexto, valor):
    ruta = crear_planeacion_xlsx(contexto.uploads.name)
    libro = openpyxl.load_workbook(ruta)
    libro['PLANEACION']['I18'] = valor
    libro.save(ruta)
    libro.close()
    with pytest.raises(ValueError, match='finito'):
        parsear_planeacion(ruta)


def test_curva_con_un_corte_no_inventa_un_hito_anterior(contexto):
    fecha = datetime(2026, 10, 21, 8)
    db.session.add(PuntajeHistorico(
        ficha_id=contexto.ficha.id, aprendiz_id=contexto.aprendiz.id,
        fecha_corte=fecha, puntaje_total=80, puntaje_asistencia=75,
        puntaje_evidencias=85, puntaje_juicios=0, posicion=2,
    ))
    db.session.commit()
    curva = obtener_curva_rendimiento_aprendiz(contexto.ficha.id, contexto.aprendiz.id, hoy=fecha.date())
    assert len(curva['puntos']) == 1
    assert curva['puntos'][0]['fecha'] == fecha.date()
    assert curva['mejor_puntaje'] == curva['peor_puntaje'] == 80
    assert curva['diferencia_ultimo'] == 0
    assert curva['tendencia'] == 'estable'


def test_panorama_conserva_spi_cero_calculado():
    desempeno = _calcular_desempeno_ficha({}, {}, {'evm': {'spi': 0.0, 'label': 'Sin avance'}}, {}, {})
    assert desempeno['spi'] == 0.0
    assert desempeno['spi_label'] == 'Sin avance'


@pytest.mark.parametrize(('tipo', 'url'), [
    ('github', 'https://example.com/github.com/proyecto'),
    ('github', 'https://github.com.example.com/proyecto'),
    ('github', 'https://example.com/?destino=github.com'),
    ('notion', 'https://example.com/notion.so/bitacora'),
    ('notion', 'https://notion.so.example.com/bitacora'),
    ('notion', 'https://notion.net/bitacora'),
])
def test_portafolio_rechaza_dominio_ajeno_aunque_la_ruta_contenga_la_marca(tipo, url):
    with pytest.raises(ValueError, match='dominio'):
        normalizar_url_portafolio(url, tipo=tipo)


@pytest.mark.parametrize(('tipo', 'url'), [
    ('github', 'github.com/aprendiz/proyecto'),
    ('github', 'https://www.github.com/aprendiz/proyecto'),
    ('notion', 'https://aprendiz.notion.site/Bitacora'),
    ('notion', 'notion.so/aprendiz/bitacora'),
])
def test_portafolio_conserva_dominios_y_subdominios_permitidos(tipo, url):
    esperado = url if url.startswith('https://') else 'https://' + url
    assert normalizar_url_portafolio(url, tipo=tipo) == esperado


def test_archivar_y_crear_grupos_puede_revertirse_como_una_sola_operacion(contexto):
    anterior = Grupo(ficha_id=contexto.ficha.id, nombre='Grupo anterior', activo=True)
    db.session.add(anterior)
    db.session.commit()
    anterior_id = anterior.id
    archivar_grupos_de_ficha(contexto.ficha.id, confirmar=False)
    nuevos = crear_grupos_aleatorios(contexto.ficha.id, [contexto.aprendiz.id], 2, confirmar=False)
    assert len(nuevos) == 1
    assert nuevos[0].id != anterior_id
    db.session.rollback()
    assert Grupo.query.count() == 1
    assert db.session.get(Grupo, anterior_id).activo is True


def test_estado_de_turno_puede_revertirse_junto_con_sus_efectos(contexto):
    turno = AtencionService.solicitar_turno(contexto.ficha.id, contexto.aprendiz.id, 'Orientación')
    turno_id = turno.id
    actualizado = AtencionService.cambiar_estado(turno.id, 'atendido', confirmar=False)
    assert actualizado.estado == 'atendido'
    assert actualizado.completado_en is not None
    db.session.rollback()
    from app.models.atencion import TurnoAtencion
    persistido = db.session.get(TurnoAtencion, turno_id)
    assert persistido.estado == 'esperando'
    assert persistido.completado_en is None


@pytest.mark.parametrize('tipo', ['planeacion', 'programa'])
def test_recuperacion_de_documento_no_acepta_un_codigo_de_programa_distinto(contexto, tipo):
    contexto.ficha.codigo_programa = '999999'
    contexto.ficha.nombre_programa = 'Tecnólogo en Analisis y Desarrollo de Software'
    db.session.commit()
    if tipo == 'planeacion':
        crear_planeacion_xlsx(contexto.uploads.name)
        buscar = buscar_archivo_planeacion_original
    else:
        crear_programa_pdf(contexto.uploads.name)
        buscar = buscar_archivo_programa_original
    assert buscar(contexto.ficha) is None


@pytest.mark.parametrize('tipo', ['planeacion', 'programa'])
def test_recuperacion_reconoce_los_metadatos_del_parser_real(contexto, tipo):
    contexto.ficha.codigo_programa = '228118'
    contexto.ficha.nombre_programa = 'Tecnólogo en Analisis y Desarrollo de Software'
    db.session.commit()
    if tipo == 'planeacion':
        ruta = crear_planeacion_xlsx(contexto.uploads.name)
        buscar = buscar_archivo_planeacion_original
    else:
        ruta = crear_programa_pdf(contexto.uploads.name)
        buscar = buscar_archivo_programa_original
    encontrado = buscar(contexto.ficha)
    assert encontrado is not None
    assert encontrado['path'] == ruta.resolve()
    assert encontrado['metadata']['codigo_programa'].startswith('228118')


def test_recuperacion_informa_fuentes_nuevas_y_cola_no_disponible(contexto):
    db.session.add(JuicioEvaluativo(
        ficha_id=contexto.ficha.id, aprendiz_id=contexto.aprendiz.id,
        competencia='Construir software', resultado_aprendizaje='RAP 01', juicio='APROBADO',
    ))
    db.session.commit()
    with patch('app.services.importacion_jobs.encolar_importacion', side_effect=ColaImportacionesNoDisponible('cola no disponible')):
        resultado = recuperar_fuentes_ficha(contexto.ficha, contexto.instructor.id)
    assert resultado['ok'] is True
    assert resultado['nuevas'] == ['Reporte de Juicios']
    assert resultado['total_nuevas'] == 1
    assert resultado['recalculo_pendiente'] is True
    assert resultado['advertencias']
    trabajo = ImportacionJob.query.one()
    assert trabajo.estado == 'error' and trabajo.error == 'cola no disponible'
    assert (Path(contexto.uploads.name) / resultado['reporte'].ruta_archivo).is_file()


@pytest.mark.parametrize('campo', ['semanas_racha', 'horas_entrega_anticipada'])
def test_ranking_con_ventanas_grandes_no_desborda_el_calendario(contexto, campo):
    config = ConfiguracionRanking(ficha_id=contexto.ficha.id)
    setattr(config, campo, 2 ** 31 - 1)
    db.session.add(config)
    ahora = datetime(2026, 10, 21, 12)
    tarea = Tarea(ficha_id=contexto.ficha.id, instructor_id=contexto.instructor.id,
                  titulo='Evidencia verificable', creada_en=ahora - timedelta(days=1),
                  fecha_limite=ahora + timedelta(days=2))
    db.session.add(tarea)
    db.session.flush()
    db.session.add(Entrega(tarea_id=tarea.id, aprendiz_id=contexto.aprendiz.id,
                          fecha_entrega=ahora, enlace_repositorio='https://example.com/evidencia'))
    db.session.commit()
    filas, _ = calcular_ranking(contexto.ficha.id, ahora=ahora)
    assert len(filas) == 1
    assert filas[0]['aprendiz'].id == contexto.aprendiz.id
    assert filas[0]['tiene_racha'] is False
    if campo == 'horas_entrega_anticipada':
        assert filas[0]['bonus'] == 0


def test_revision_automatica_revierte_una_ficha_fallida_y_continua_con_la_siguiente(contexto):
    segunda = Ficha(codigo='4000002', codigo_ficha='4000002', nombre_programa='ADSO',
                    instructor_id=contexto.instructor.id,
                    fecha_inicio=contexto.ficha.fecha_inicio, fecha_fin=contexto.ficha.fecha_fin)
    db.session.add(segunda)
    db.session.commit()
    segunda_id = segunda.id
    ficha_fallida_id = contexto.ficha.id
    db.session.execute(text(
        'CREATE TRIGGER impedir_configuracion BEFORE INSERT ON configuracion_alertas_comite '
        f'WHEN NEW.ficha_id = {ficha_fallida_id} BEGIN SELECT RAISE(ABORT, \'fallo controlado de almacenamiento\'); END'
    ))
    db.session.commit()
    assert ejecutar_revision_automatica() == 1
    from app.models.alertas import ConfiguracionAlertasComite
    assert ConfiguracionAlertasComite.query.filter_by(ficha_id=ficha_fallida_id).count() == 0
    assert ConfiguracionAlertasComite.query.filter_by(ficha_id=segunda_id).count() == 1


@pytest.mark.parametrize('tipo', ['planeacion', 'programa'])
def test_recuperacion_fallida_no_deja_copias_sin_registro(contexto, tipo):
    contexto.ficha.codigo_programa = '228118'
    contexto.ficha.nombre_programa = 'Tecnólogo en Analisis y Desarrollo de Software'
    db.session.commit()
    if tipo == 'planeacion':
        ruta = crear_planeacion_xlsx(contexto.uploads.name)
        recuperar = asegurar_version_planeacion
    else:
        ruta = crear_programa_pdf(contexto.uploads.name)
        recuperar = asegurar_version_programa
    original = ruta.read_bytes()
    db.session.execute(text(
        'CREATE TRIGGER impedir_actualizacion_version BEFORE UPDATE ON archivos_ficha_versiones '
        "BEGIN SELECT RAISE(ABORT, 'fallo controlado de almacenamiento'); END"
    ))
    db.session.commit()
    assert recuperar(contexto.ficha) is None
    assert ruta.read_bytes() == original
    assert [p for p in Path(contexto.uploads.name).rglob('*') if p.is_file()] == [ruta]


def test_reporte_reconstruido_fallido_no_deja_archivo_sin_registro(contexto):
    db.session.add(JuicioEvaluativo(
        ficha_id=contexto.ficha.id, aprendiz_id=contexto.aprendiz.id,
        competencia='Construir software', resultado_aprendizaje='RAP 01', juicio='APROBADO',
    ))
    db.session.commit()
    db.session.execute(text(
        'CREATE TRIGGER impedir_insercion_version BEFORE INSERT ON archivos_ficha_versiones '
        "BEGIN SELECT RAISE(ABORT, 'fallo controlado de almacenamiento'); END"
    ))
    db.session.commit()
    with pytest.raises(IntegrityError):
        asegurar_version_reporte(contexto.ficha)
    assert [p for p in Path(contexto.uploads.name).rglob('*') if p.is_file()] == []

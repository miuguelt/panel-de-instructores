"""Lectura de archivos y lotes sin duplicar su contenido completo en memoria."""

from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
from werkzeug.datastructures import FileStorage

from app import db
from app.models import Aprendiz, JuicioEvaluativo
from app.services import importacion_ficha as servicio
from app.services.planeacion import parsear_planeacion
from app.services import programa_formacion
from tests import test_importacion_ficha as soporte
from tests.apoyo_planeacion import crear_planeacion_xlsx
from tests.apoyo_programa import crear_programa_pdf


class ArchivoSinLecturaCompleta(FileStorage):
    """Impide que la aplicación duplique todo el adjunto mediante read()."""

    def read(self, cantidad=-1):
        assert cantidad >= 0, 'La aplicación no debe leer el archivo completo.'
        return self.stream.read(cantidad)


def reporte(filas):
    libro = openpyxl.Workbook(write_only=True)
    hoja = libro.create_sheet()
    hoja.append(['Documento', 'Nombre', 'Competencia', 'Resultado', 'Juicio'])
    for fila in filas:
        hoja.append(fila)
    stream = BytesIO()
    libro.save(stream)
    libro.close()
    stream.seek(0)
    return ArchivoSinLecturaCompleta(stream=stream, filename='reporte.xlsx')


def test_lector_xlsx_no_lee_adjunto_completo_y_registros_se_recorren_dos_veces():
    archivo = reporte(([str(i), 'Ana', 'Competencia', f'RAP {i}', 'APROBADO'] for i in range(2100)))
    metadata, registros = servicio._leer_archivo(archivo)
    assert metadata == {}
    assert not isinstance(registros, list)
    assert len(registros) == 2100
    assert registros[0]['documento'] == '0'
    assert registros[-1]['documento'] == '2099'
    assert sum(1 for _ in registros) == 2100
    assert sum(1 for _ in registros) == 2100
    registros.close()


def test_xls_real_desde_ruta_no_requiere_duplicar_bytes():
    ruta = Path(__file__).parent / 'fixtures/rendimiento_metadatos.xls'
    with ruta.open('rb') as stream:
        archivo = ArchivoSinLecturaCompleta(stream=stream, filename=ruta.name)
        metadata, registros = servicio._leer_archivo(archivo)
    assert metadata['codigo_ficha'] == '3999999'
    assert len(registros) == 35
    assert registros[-1]['documento'] == '34'
    registros.close()


def test_planeacion_usa_lectura_iterativa_y_conserva_combinadas(tmp_path):
    ruta = crear_planeacion_xlsx(tmp_path)
    cargar = openpyxl.load_workbook
    with patch('app.services.planeacion.openpyxl.load_workbook', wraps=cargar) as apertura:
        contenido = parsear_planeacion(ruta)
    assert apertura.call_args.kwargs['read_only'] is True
    assert contenido['resumen']['horas_total'] == 47
    assert contenido['unidades'][0]['fase'] == 'ANÁLISIS'
    assert contenido['unidades'][0]['filas_planificacion'] == 2
    assert contenido['metadata']['codigo_programa'] == '228118 v 1.0'


def test_pdf_se_abre_como_stream_y_cierra_archivo_al_terminar(tmp_path):
    ruta = crear_programa_pdf(tmp_path)
    lector = programa_formacion.PdfReader
    with patch.object(programa_formacion, 'PdfReader', wraps=lector) as apertura:
        programa = programa_formacion.parsear_programa(ruta)
    stream = apertura.call_args.args[0]
    assert hasattr(stream, 'read'), 'El lector debe recibir un stream, no duplicar el PDF desde su ruta.'
    assert stream.closed
    assert programa['metadata']['codigo_programa'] == '228118'
    assert programa['resumen']['competencias'] == 3


def test_secuencia_temporal_rechaza_indices_fuera_de_rango_y_se_cierra():
    registros = servicio._RegistrosDisco([{'numero': 1}, {'numero': 2}])
    assert registros[-2] == {'numero': 1}
    for indice in (-3, 2):
        with pytest.raises(IndexError):
            registros[indice]
    registros.close()
    assert registros.archivo.closed


def test_spool_limpia_temporal_si_el_iterador_falla():
    def filas():
        yield {'numero': 1}
        raise ValueError('Archivo incompleto')

    temporales = []
    crear = servicio.tempfile.TemporaryFile

    def registrar(*args, **kwargs):
        temporal = crear(*args, **kwargs)
        temporales.append(temporal)
        return temporal

    with patch.object(servicio.tempfile, 'TemporaryFile', registrar):
        with pytest.raises(ValueError, match='Archivo incompleto'):
            servicio._RegistrosDisco(filas())
    assert len(temporales) == 1
    assert temporales[0].closed


@pytest.fixture
def contexto_importacion():
    caso = soporte.ImportacionFichaTestCase()
    caso.setUp()
    try:
        yield caso
    finally:
        caso.tearDown()


def test_lotes_conservan_ultima_fila_y_contadores_sin_commit(contexto_importacion):
    caso = contexto_importacion
    filas = [[str(i), 'Ana', 'Competencia', f'RAP {i}', 'POR EVALUAR'] for i in range(1001)]
    filas.append(['0', 'Ana María', 'Competencia', 'RAP 0', 'APROBADO'])
    resultado = servicio.importar_archivo(reporte(filas), caso.ficha, caso.a.id)
    assert resultado['nuevos'] == 1001
    assert resultado['juicios_nuevos'] == 1001
    assert resultado['juicios_repetidos'] == 1
    assert resultado['juicios_actualizados'] == 0
    assert Aprendiz.query.filter_by(documento='0').one().nombre == 'Ana María'
    assert JuicioEvaluativo.query.filter_by(resultado_aprendizaje='RAP 0').one().juicio == 'APROBADO'
    db.session.rollback()
    assert Aprendiz.query.count() == 0
    assert JuicioEvaluativo.query.count() == 0
